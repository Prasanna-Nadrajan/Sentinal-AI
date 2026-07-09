import docx

def create_doc(txt_file, doc_file):
    doc = docx.Document()
    with open(txt_file, 'r', encoding='utf-8') as f:
        text = f.read()
    
    for paragraph in text.split('\n\n'):
        if paragraph.strip() == '---':
            doc.add_paragraph('--------------------------------------------------')
        else:
            doc.add_paragraph(paragraph.strip())
            
    doc.save(doc_file)

create_doc("sentinel_demo_script.txt", "sentinel_demo_script.docx")
create_doc("sentinel_showcase_script.txt", "sentinel_showcase_script.docx")
